# 카카오 AdFit과 쿠팡 파트너스 실행 절차서를 생성하는 스크립트
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(r"C:\Users\mjs5n\Desktop\길드원 따라가기 광고 플랫폼 신청 및 적용 절차.docx")
NAVY = "172235"
LIGHT_BLUE = "EAF0F7"
PALE_ORANGE = "FFF1E6"
ORANGE = RGBColor(232, 111, 48)
GRAY = RGBColor(92, 103, 117)


def set_run_font(run, name="Malgun Gothic", size=10.5, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    keep_with_next(p)
    run = p.add_run(text)
    set_run_font(run, size=15 if level == 1 else 12, bold=True, color=RGBColor(0, 0, 0))
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_check(doc, text, owner=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run("□ ")
    set_run_font(r, size=11, bold=True, color=ORANGE)
    if owner:
        tag = p.add_run(f"[{owner}] ")
        set_run_font(tag, size=9.5, bold=True, color=GRAY)
    body = p.add_run(text)
    set_run_font(body)
    return p


def add_numbered(doc, number, title, detail=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.65)
    p.paragraph_format.first_line_indent = Cm(-0.65)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    n = p.add_run(f"{number}. ")
    set_run_font(n, size=10.5, bold=True, color=ORANGE)
    t = p.add_run(title)
    set_run_font(t, bold=True)
    if detail:
        d = p.add_run(f"  {detail}")
        set_run_font(d, color=GRAY)
    return p


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relation_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "C85B28")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    r_fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia"):
        r_fonts.set(qn(f"w:{attr}"), "Malgun Gothic")
    r_pr.append(r_fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    r_pr.append(size)
    run.append(r_pr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_source(doc, label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("• ")
    set_run_font(r, color=ORANGE)
    add_hyperlink(p, label, url)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.7)
section.left_margin = Cm(2.1)
section.right_margin = Cm(2.1)

styles = doc.styles
for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
    style = styles[style_name]
    style.font.name = "Malgun Gothic"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    style.font.color.rgb = RGBColor(0, 0, 0)

title_ppr = styles["Title"]._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.paragraph_format.space_after = Pt(7)
run = title.add_run("길드원 따라가기 광고 플랫폼 신청 및 적용 절차")
set_run_font(run, size=20.5, bold=True, color=RGBColor(0, 0, 0))

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(15)
r = subtitle.add_run("카카오 AdFit과 쿠팡 파트너스 운영 실행서")
set_run_font(r, size=11.5, bold=True, color=ORANGE)
r2 = subtitle.add_run("   |   기준일 2026년 9월 16일")
set_run_font(r2, size=9.5, color=GRAY)

add_body(doc, "이 문서는 guildfollow.com에 광고와 제휴 링크를 적용하기 위해 서비스 소유자가 직접 해야 할 신청 절차와 개발 작업을 구분한 실행 체크리스트입니다. 현재 조건에서는 쿠팡 파트너스를 먼저 적용하고, 카카오 AdFit은 제휴 문의와 사전 승인을 받은 뒤 적용하는 순서가 가장 현실적입니다.")
add_body(doc, "정책과 정산 조건은 변경될 수 있으므로 실제 신청 시 각 플랫폼의 최신 안내를 다시 확인합니다. 계정 비밀번호, 인증번호, API 키는 개발 작업을 위해 공유하지 않습니다.")

add_heading(doc, "1 실행 순서 요약")
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Cm(1.3), Cm(3.2), Cm(6.6), Cm(4.2)]
headers = ["순서", "플랫폼", "핵심 행동", "완료 기준"]
for i, cell in enumerate(table.rows[0].cells):
    cell.width = widths[i]
    shade_cell(cell, NAVY)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p.add_run(headers[i])
    set_run_font(rr, size=9.5, bold=True, color=RGBColor(255, 255, 255))
rows = [
    ("1", "쿠팡 파트너스", "가입하고 guildfollow.com을 활동 페이지로 등록", "파트너 ID와 링크 생성 가능"),
    ("2", "쿠팡 파트너스", "게이밍 관련 상품 링크 3개에서 5개 생성", "상품명과 제휴 링크 확보"),
    ("3", "웹 적용", "경제적 이해관계 표시와 함께 추천 카드 배치", "운영 화면과 모바일 확인"),
    ("4", "카카오 AdFit", "공식 페이지에서 제휴 문의 접수", "인증 코드 또는 후속 안내 수신"),
    ("5", "카카오 AdFit", "매체와 광고 단위를 등록하고 공식 스크립트 설치", "매체 심사 승인"),
]
for row_index, values in enumerate(rows, start=1):
    cells = table.add_row().cells
    for i, value in enumerate(values):
        cells[i].width = widths[i]
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cells[i])
        if row_index % 2 == 0:
            shade_cell(cells[i], "F5F7FA")
        p = cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 1) else WD_ALIGN_PARAGRAPH.LEFT
        rr = p.add_run(value)
        set_run_font(rr, size=9.2, bold=(i == 1))
set_table_borders(table)

add_heading(doc, "2 쿠팡 파트너스 신청과 적용")
add_body(doc, "쿠팡 파트너스는 배너 노출 자체보다 제휴 링크를 통해 구매가 발생했을 때 수수료를 받는 구조입니다. 외부 광고 스크립트를 웹 전체에 설치하지 않고 자체 추천 카드로 표현할 수 있어 현재 서비스에 먼저 적용하기 좋습니다.")

add_heading(doc, "2.1 서비스 소유자가 직접 할 일", level=2)
add_numbered(doc, 1, "쿠팡 파트너스 가입", "본인 쿠팡 계정으로 로그인하고 파트너 가입을 완료합니다.")
add_numbered(doc, 2, "활동 페이지 등록", "웹사이트 주소에 https://guildfollow.com 을 입력합니다.")
add_numbered(doc, 3, "본인 정보 입력", "요청되는 연락처와 본인 인증 정보를 입력합니다.")
add_numbered(doc, 4, "상품 링크 생성", "게이밍 키보드, 마우스, 모니터, 헤드셋 등 서비스 이용자와 관련 있는 상품을 우선합니다.")
add_numbered(doc, 5, "개발 작업용 자료 전달", "상품명, 간단한 추천 문구, 파트너스 링크만 전달합니다. 계정 비밀번호는 전달하지 않습니다.")
add_numbered(doc, 6, "배치 화면 제출", "파트너스 화면에서 활동 페이지 캡처나 최종 심사를 요구하면 실제 적용 화면을 제출합니다.")
add_numbered(doc, 7, "정산 정보 등록", "정산 메뉴가 활성화되면 본인 명의 계좌와 필요한 세금 정보를 입력합니다.")

add_heading(doc, "2.2 상품 링크 준비 목록", level=2)
for text in (
    "게이밍 키보드 1개",
    "게이밍 마우스 또는 마우스패드 1개",
    "모니터 또는 모니터암 1개",
    "헤드셋 또는 스피커 1개",
    "장시간 플레이용 의자나 손목 받침대 1개",
):
    add_check(doc, text, "사용자")

add_heading(doc, "2.3 개발 작업과 게시 확인", level=2)
for text, owner in (
    ("상품 링크를 설정 파일에 등록하고 비밀번호나 계정 정보를 저장하지 않습니다.", "개발"),
    ("PC와 모바일에서 화면을 가리지 않는 게이밍 장비 추천 카드로 표시합니다.", "개발"),
    ("광고 제휴 링크 표기와 수수료 발생 가능성을 카드 가까이에 명확히 표시합니다.", "개발"),
    ("API 키 입력, 설정 저장, 새로고침 조작부 주변에는 광고를 두지 않습니다.", "개발"),
    ("적용 후 guildfollow.com 실제 화면을 확인하고 필요한 캡처를 준비합니다.", "공동"),
):
    add_check(doc, text, owner)

add_body(doc, "권장 공개 문구. ‘광고·제휴 링크’와 ‘이 링크를 통해 구매하면 서비스 운영자에게 일정액의 수수료가 지급될 수 있습니다’를 함께 표시합니다.", bold_lead="권장 공개 문구.")
add_body(doc, "정산 참고. 공식 가이드 기준으로 가입 후 링크 생성이 가능하며, 최종 승인과 전월까지의 누적 수익 조건을 충족하면 정산정보 입력이 활성화됩니다. 실제 금액과 조건은 파트너스 대시보드의 최신 안내를 우선합니다.", bold_lead="정산 참고.")

add_heading(doc, "3 카카오 AdFit 신청과 적용")
add_body(doc, "카카오 AdFit은 현재 누구나 즉시 매체를 등록하는 공개 가입형 절차가 아닙니다. 먼저 공식 페이지에서 제휴 문의를 접수하고, 카카오와의 협의 후 인증 코드나 후속 안내를 받아야 가입과 매체 등록을 진행할 수 있습니다.")

add_heading(doc, "3.1 제휴 문의에 준비할 내용", level=2)
items = [
    ("서비스명", "길드원 따라가기"),
    ("서비스 주소", "https://guildfollow.com"),
    ("서비스 설명", "메이플스토리 공식 API를 이용한 캐릭터 성장 기록과 길드 비교 서비스"),
    ("이용 방식", "로그인 없이 주요 조회 기능을 사용하며 사용자가 본인의 NEXON 서비스 키로 직접 조회"),
    ("지원 환경", "PC와 모바일 웹"),
    ("광고 예정 위치", "데스크톱 좌우 레일 또는 콘텐츠 하단의 사용자 조작을 가리지 않는 영역"),
    ("운영 계획", "정식 도메인을 이용해 지속 운영하며 개인정보 안내와 이용 안내를 제공"),
]
info = doc.add_table(rows=1, cols=2)
info.alignment = WD_TABLE_ALIGNMENT.CENTER
info.autofit = False
info.rows[0].cells[0].width = Cm(3.5)
info.rows[0].cells[1].width = Cm(11.8)
for i, label in enumerate(("항목", "기재 내용")):
    shade_cell(info.rows[0].cells[i], NAVY)
    set_cell_margins(info.rows[0].cells[i])
    p = info.rows[0].cells[i].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = p.add_run(label)
    set_run_font(rr, size=9.5, bold=True, color=RGBColor(255, 255, 255))
for index, (label, value) in enumerate(items, start=1):
    cells = info.add_row().cells
    for c in cells:
        set_cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if index % 2 == 0:
        shade_cell(cells[0], "F5F7FA")
        shade_cell(cells[1], "F5F7FA")
    rr = cells[0].paragraphs[0].add_run(label)
    set_run_font(rr, size=9.2, bold=True)
    rr = cells[1].paragraphs[0].add_run(value)
    set_run_font(rr, size=9.2)
set_table_borders(info)

add_heading(doc, "3.2 승인 이후 절차", level=2)
add_numbered(doc, 1, "인증 코드로 가입", "카카오가 전달한 안내에 따라 AdFit 계정을 만듭니다.")
add_numbered(doc, 2, "회원 유형과 본인 인증", "개인 또는 사업자 유형을 정확하게 선택합니다.")
add_numbered(doc, 3, "웹 매체 등록", "매체 주소에 https://guildfollow.com 을 등록합니다.")
add_numbered(doc, 4, "광고 단위 생성", "승인받은 배치에 맞춰 광고 단위를 생성합니다.")
add_numbered(doc, 5, "설치 정보 전달", "광고 단위 ID나 공식 설치 스크립트만 개발 작업에 사용합니다.")
add_numbered(doc, 6, "시험 배포", "최신 공식 스크립트가 실제로 광고 요청을 보내도록 제한된 위치에 설치합니다.")
add_numbered(doc, 7, "매체 심사 신청", "콘텐츠와 광고 배치, 정상 광고 호출 상태를 확인받습니다.")
add_numbered(doc, 8, "정산 계좌 등록", "승인 후 회원 유형에 맞는 계좌와 세금 자료를 준비합니다.")

add_heading(doc, "3.3 심사 전 점검", level=2)
for text, owner in (
    ("서비스 주소와 주요 화면이 외부에서 정상적으로 열립니다.", "공동"),
    ("개인정보 안내와 이용 안내, 운영자 연락 경로를 확인합니다.", "개발"),
    ("콘텐츠가 임시 화면이나 빈 화면이 아니며 주요 기능을 사용할 수 있습니다.", "공동"),
    ("광고가 버튼, API 키 입력란, 메뉴와 겹치지 않습니다.", "개발"),
    ("모바일과 PC에서 레이아웃이 깨지지 않습니다.", "개발"),
    ("공식 스크립트가 최신 버전이며 실제 광고 요청이 발생합니다.", "개발"),
):
    add_check(doc, text, owner)

add_body(doc, "정산 참고. 현재 공식 정책은 지급 신청 최소액을 5만 원으로 안내합니다. 개인 회원은 원천징수 대상이 될 수 있고, 등록 계좌의 예금주와 회원 정보가 일치해야 합니다. 사업자는 세금계산서 등 추가 서류가 필요할 수 있습니다.", bold_lead="정산 참고.")

add_heading(doc, "4 자료 전달 원칙")
add_check(doc, "전달 가능. 상품명, 추천 문구, 쿠팡 파트너스 링크, AdFit 광고 단위 ID, 공개 설치 스크립트")
add_check(doc, "전달 금지. 계정 비밀번호, 문자 인증번호, 계좌 비밀번호, NEXON API 키")
add_check(doc, "신분증과 사업자 서류는 각 플랫폼의 공식 입력 화면에만 제출합니다.")
add_check(doc, "정산 계좌는 각 플랫폼의 공식 대시보드에서 직접 등록합니다.")

add_heading(doc, "5 최종 실행 체크리스트")
checklist = [
    ("쿠팡 파트너스 가입 완료", "사용자"),
    ("활동 페이지 guildfollow.com 등록", "사용자"),
    ("상품 링크 3개에서 5개 생성", "사용자"),
    ("상품명과 링크 전달", "사용자"),
    ("추천 카드와 경제적 이해관계 문구 구현", "개발"),
    ("PC와 모바일 실제 화면 점검", "공동"),
    ("쿠팡 최종 심사 자료 제출", "사용자"),
    ("카카오 AdFit 제휴 문의 접수", "사용자"),
    ("카카오 인증 코드 또는 승인 안내 수신", "사용자"),
    ("AdFit 매체와 광고 단위 등록", "사용자"),
    ("공식 광고 스크립트 설치와 시험 배포", "개발"),
    ("AdFit 매체 심사 신청", "사용자"),
    ("승인 후 정산 계좌와 세금 정보 등록", "사용자"),
]
for text, owner in checklist:
    add_check(doc, text, owner)

add_heading(doc, "6 공식 확인 링크")
add_source(doc, "쿠팡 파트너스", "https://partners.coupang.com/")
add_source(doc, "쿠팡 파트너스 공식 이용 가이드", "https://partners.coupangcdn.com/partners-guide/partners-guide-20250714121952.pdf")
add_source(doc, "카카오 AdFit 공식 안내", "https://adfit.kakao.com/info")
add_source(doc, "카카오 AdFit 운영정책", "https://adfit.kakao.com/web/html/use_kakao.html")
add_source(doc, "공정거래위원회 추천 보증 관련 안내", "https://www.ftc.go.kr/www/selectBbsNttView.do?bordCd=3&key=12&nttSn=43669&pageIndex=2&pageUnit=10&rltnNttSn=46006&searchCnd=all&searchCtgry=01%2C02&searchViolt=0609")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("길드원 따라가기 광고 플랫폼 운영 절차")
set_run_font(fr, size=8.5, color=GRAY)

for sec in doc.sections:
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.7)
    sec.left_margin = Cm(2.1)
    sec.right_margin = Cm(2.1)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
